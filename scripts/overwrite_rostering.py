#!/usr/bin/env python3
"""Overwrite Tema_Kumasi_Rostering_PERCLOS.xlsx from a coded-night logs/ folder."""
from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from fatigue_wa.pairing import pair_logs


def read(path: Path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def f(row, key):
    try:
        return float(row.get(key) or 0)
    except (TypeError, ValueError):
        return 0.0


def gold_rule(labels):
    checking = [r for r in labels if r.get("paired") == "yes" and r.get("label") == "checking"]
    gone = [r for r in labels if r.get("label") == "gone"]
    return checking, gone


def roc_point(pos, neg, thr):
    pos, neg = np.asarray(pos), np.asarray(neg)
    if pos.size == 0 or neg.size == 0:
        return None, None
    return float((pos >= thr).mean()), float((neg >= thr).mean())


def p95_table(hz, stops):
    if not hz:
        return []
    t0 = min(f(r, "unix_time") for r in hz)
    stop_t = [f(s, "unix_time") for s in stops]
    stop_stim = {f(s, "unix_time"): (s.get("stimulant") or "none") for s in stops}

    def stim_at(t):
        prior = [st for st in stop_t if st <= t]
        if not prior:
            return "none"
        return stop_stim[max(prior)]

    buckets = defaultdict(list)
    for r in hz:
        t = f(r, "unix_time")
        buckets[(int((t - t0) // 3600), stim_at(t))].append(f(r, "perclos"))
    hours = sorted({h for h, _ in buckets})
    rows = []
    for h in hours:
        none = buckets.get((h, "none"), [])
        ata = buckets.get((h, "ataya"), [])
        def p95(v):
            return float(np.percentile(v, 95)) if v else None
        pn, pa = p95(none), p95(ata)
        peak = max([x for x in (pn, pa) if x is not None], default=None)
        if peak is None:
            flag = "no data"
        elif peak >= 0.30:
            flag = "SPLIT or STOP"
        elif peak >= 0.22:
            flag = "shorten next duty"
        else:
            flag = "within band"
        rows.append((h, pn, pa, len(none) + len(ata), flag))
    return rows


def startle(hz, misses, window=30.0, floor=0.45):
    out = []
    if not hz or not misses:
        return out
    series = []
    for r in hz:
        conf = f(r, "mesh_conf") if r.get("mesh_conf") else (1.0 if r.get("face") not in ("0", "false", "") else 0.0)
        series.append((f(r, "unix_time"), conf))
    for i, m in enumerate(misses, 1):
        t0 = f(m, "unix_time")
        w = [c for t, c in series if t0 <= t <= t0 + window]
        out.append((i, t0, float(np.mean([c < floor for c in w])) if w else None, len(w)))
    return out


def write_xlsx(path: Path, summary, p95, depot, trunk, clips, banner: str) -> None:
    thin = Border(left=Side(style="thin", color="D0D5DD"), right=Side(style="thin", color="D0D5DD"),
                  top=Side(style="thin", color="D0D5DD"), bottom=Side(style="thin", color="D0D5DD"))
    hdr_f = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    title = Font(name="Calibri", bold=True, size=16, color="1F4E79")
    sec = Font(name="Calibri", bold=True, size=12, color="1F4E79")
    blue = Font(name="Calibri", color="0000FF")
    grey = Font(name="Calibri", italic=True, color="666666", size=9)
    wait = PatternFill("solid", fgColor="FFF3CD")
    good = PatternFill("solid", fgColor="D4EDDA")
    bad = PatternFill("solid", fgColor="F8D7DA")

    def head(ws, r, vals):
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.fill = hdr_f
            c.font = hdr_font
            c.border = thin

    wb = Workbook()
    ws = wb.active
    ws.title = "Haul cover"
    ws["A1"] = "Tema-Kumasi coded night — rostering workbook"
    ws["A1"].font = title
    ws.merge_cells("A1:F1")
    ws["A2"] = banner
    ws["A2"].font = grey
    status = "CODED" if summary["alerts"] else "WAITING FOR FIELD CSVs"
    rows = [
        ("Status", status),
        ("Banners", summary["alerts"]),
        ("Labels", summary["labels"]),
        ("checking gold", summary["checking"]),
        ("gone", summary["gone"]),
        ("gone rate", summary["gone_rate"]),
        ("Gold rule", "paired=yes AND label=checking"),
        ("Names", "none"),
    ]
    head(ws, 4, ["Field", "Value"])
    for i, (a, b) in enumerate(rows, 5):
        ws.cell(i, 1, a).border = thin
        c = ws.cell(i, 2, b)
        c.border = thin
        c.font = blue
        if a == "gone rate" and isinstance(b, float):
            c.number_format = "0.0%"
            c.fill = wait
        if a == "Status":
            c.fill = good if status == "CODED" else wait

    ws2 = wb.create_sheet("PERCLOS p95 rostering")
    ws2["A1"] = "PERCLOS 95th percentile by haul hour x stimulant"
    ws2["A1"].font = title
    head(ws2, 4, ["Haul hour", "p95 none", "p95 ataya", "n samples", "Roster flag"])
    if not p95:
        ws2["A5"] = "No perclos_1hz.csv yet"
        ws2["A5"].fill = wait
    for i, (h, pn, pa, n, flag) in enumerate(p95, 5):
        ws2.cell(i, 1, h).border = thin
        for col, val in ((2, pn), (3, pa)):
            cell = ws2.cell(i, col, val if val is not None else "—")
            cell.border = thin
            if isinstance(val, float):
                cell.number_format = "0.0%"
        ws2.cell(i, 4, n).border = thin
        fl = ws2.cell(i, 5, flag)
        fl.border = thin
        fl.fill = bad if flag == "SPLIT or STOP" else wait if flag != "within band" else good

    ws3 = wb.create_sheet("Gold set and operating points")
    ws3["A1"] = "Sensitivity only on paired=yes AND label=checking"
    ws3["A1"].font = title
    head(ws3, 4, ["Item", "N", "Rate", "Rule"])
    items = [
        ("Detector banners", summary["alerts"], 1.0 if summary["alerts"] else 0, "alerts.csv"),
        ("checking (gold)", summary["checking"], summary["checking"] / summary["alerts"] if summary["alerts"] else 0, "paired=yes and checking"),
        ("gone", summary["gone"], summary["gone_rate"], "discard from gold"),
    ]
    for i, (name, n, rate, rule) in enumerate(items, 5):
        ws3.cell(i, 1, name).border = thin
        ws3.cell(i, 2, n).border = thin
        c = ws3.cell(i, 3, rate)
        c.border = thin
        c.number_format = "0.0%"
        ws3.cell(i, 4, rule).border = thin
        if name == "gone":
            ws3.cell(i, 1).fill = wait
            c.fill = wait
    ws3["A9"] = "Operating points"
    ws3["A9"].font = sec
    head(ws3, 10, ["Point", "Cut", "TPR", "FPR"])
    for r, label, cut, pt in ((11, "Depot road", 0.45, depot), (12, "Trunk road", 0.55, trunk)):
        ws3.cell(r, 1, label).border = thin
        ws3.cell(r, 2, cut).border = thin
        ws3.cell(r, 2).font = blue
        tpr = pt[0] if pt and pt[0] is not None else "—"
        fpr = pt[1] if pt and pt[1] is not None else "—"
        c3 = ws3.cell(r, 3, tpr); c3.border = thin
        c4 = ws3.cell(r, 4, fpr); c4.border = thin
        if isinstance(tpr, float):
            c3.number_format = "0.0%"; c4.number_format = "0.0%"
    ws3["A14"] = "Startle-loss 0-30 s"
    ws3["A14"].font = sec
    head(ws3, 15, ["Clip", "Blind fraction", "Samples"])
    if not clips:
        ws3["A16"] = "No near_misses.csv yet"
        ws3["A16"].fill = wait
    for i, (idx, t0, blind, n) in enumerate(clips, 16):
        ws3.cell(i, 1, f"Near-miss {idx}").border = thin
        c = ws3.cell(i, 2, blind if blind is not None else "—")
        c.border = thin
        if isinstance(blind, float):
            c.number_format = "0.0%"
            c.fill = bad if blind >= 0.5 else wait
        ws3.cell(i, 3, n).border = thin

    for wsx in wb.worksheets:
        wsx.page_setup.fitToPage = True
        wsx.page_setup.fitToWidth = 1
        wsx.sheet_properties.pageSetUpPr.fitToPage = True
        wsx.oddHeader.left.text = "Quaye  ·  coded night  ·  no names"
        wsx.column_dimensions["A"].width = 28
        wsx.column_dimensions["B"].width = 24
        for col in range(3, 6):
            wsx.column_dimensions[get_column_letter(col)].width = 16
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--logs", default="logs")
    p.add_argument("--out", default="Tema_Kumasi_Rostering_PERCLOS.xlsx")
    args = p.parse_args()
    root = Path(args.logs)
    pair_logs(root)
    alerts = read(root / "alerts.csv")
    labels = read(root / "paired_alerts.csv")
    stops = read(root / "paired_rest_stops.csv") or read(root / "rest_stops.csv")
    hz = read(root / "perclos_1hz.csv")
    misses = read(root / "near_misses.csv")
    gold, gone = gold_rule(labels)
    pos = [f(r, "score") for r in gold if r.get("state") in ("DROWSY", "MICROSLEEP")]
    neg = [f(r, "score") for r in gold if r.get("state") == "ALERT"]
    summary = {
        "alerts": len(alerts),
        "labels": len(labels),
        "checking": len(gold),
        "gone": len(gone),
        "gone_rate": (len(gone) / len(alerts)) if alerts else 0.0,
    }
    write_xlsx(
        Path(args.out), summary, p95_table(hz, stops),
        roc_point(pos, neg, 0.45), roc_point(pos, neg, 0.55), startle(hz, misses),
        f"Overwritten from {root.resolve()}. Gold = checking only. No names.",
    )
    print("wrote", args.out)
    print("alerts", summary["alerts"], "checking", summary["checking"], "gone", summary["gone"],
          f"gone_rate {summary['gone_rate']:.1%}")


if __name__ == "__main__":
    main()
